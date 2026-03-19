#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
公式：
内核占用 = (所有 VM_KERN_* + zones) - VM_KERN_COUNT_BOOT_STOLEN + VM_KERN_COUNT_WIRED_残余
驱动占用 = sum(wired memory 中的 com.apple.*) - com.apple.filesystems.apfs
其中：
- VM_KERN_* 为 wired memory 页表中所有以 VM_KERN_ 开头的字段之和
- VM_KERN_COUNT_WIRED_残余 = total Y of X 中 X - Y 的值
- com.apple.* 仅统计 wired memory 页表中的项（maps 页表不计入）
"""

import os
import re
from collections import defaultdict
from datetime import datetime
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:
    print("错误: 需要安装 openpyxl 库")
    print("请运行: pip install openpyxl")
    exit(1)


def parse_size(text):
    text = text.strip()

    match = re.search(r'(\d+\.?\d*)([KMG])', text)
    if not match:
        return 0

    try:
        value = float(match.group(1))
    except ValueError:
        return 0

    unit = match.group(2)

    if unit == 'K':
        return int(value)
    elif unit == 'M':
        return int(value * 1024)
    elif unit == 'G':
        return int(value * 1024 * 1024)
    return 0


def parse_wired_residual(line):
    match = re.search(r'total\s+([\d.]+)([KMG])\s+of\s+([\d.]+)([KMG])', line)
    if not match:
        return 0

    used_val = float(match.group(1))
    used_unit = match.group(2)
    total_val = float(match.group(3))
    total_unit = match.group(4)

    # 转换为 KB
    used_kb = parse_size(f"{used_val}{used_unit}")
    total_kb = parse_size(f"{total_val}{total_unit}")

    return total_kb - used_kb


def parse_zprint_file(filepath):
    data = {
        'vm_kern': defaultdict(int),
        'drivers': defaultdict(int),
        'zones': 0,
        'wired_residual': 0,
        'timestamp': None,  # 文件中的时间戳
    }

    in_maps_section = False

    with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
        for line in f:
            line = line.strip()
            if not line:
                continue

            if data['timestamp'] is None:
                time_patterns = [
                    r'(\d{4}[-/]\d{2}[-/]\d{2}\s+\d{2}:\d{2}:\d{2})',
                    r'(\d{2}[-/]\d{2}[-/]\d{4}\s+\d{2}:\d{2}:\d{2})',
                    r'(\d{4}[-/]\d{2}[-/]\d{2})',
                ]
                for pattern in time_patterns:
                    match = re.search(pattern, line)
                    if match:
                        data['timestamp'] = match.group(1)
                        break

            if line.startswith('maps') and 'free' in line:
                in_maps_section = True
                continue

            if in_maps_section:
                continue

            if line.startswith('VM_KERN_'):
                parts = line.split()
                if len(parts) >= 2:
                    key = parts[0]
                    data['vm_kern'][key] += parse_size(line)

            elif line.startswith('zones '):
                data['zones'] = parse_size(line)

            elif line.startswith('total') and 'of' in line:
                residual = parse_wired_residual(line)
                if residual > 0:
                    data['wired_residual'] = residual

            elif line.startswith('com.apple.'):
                parts = line.split()
                if len(parts) >= 2:
                    key = parts[0]
                    data['drivers'][key] += parse_size(line)

    return data


def calculate_memory(data):
    vm_kern_sum = sum(data['vm_kern'].values())
    boot_stolen = data['vm_kern'].get('VM_KERN_COUNT_BOOT_STOLEN', 0)

    # 内核占用
    kernel_kb = (vm_kern_sum + data['zones']) - boot_stolen + data['wired_residual']

    # 驱动占用
    driver_sum = sum(data['drivers'].values())
    apfs_kb = data['drivers'].get('com.apple.filesystems.apfs', 0)
    driver_kb = driver_sum - apfs_kb

    return kernel_kb, driver_kb


def format_output_row(filename, data, kernel_kb, driver_kb, file_mtime):
    timestamp = data.get('timestamp')
    if not timestamp:
        timestamp = datetime.fromtimestamp(file_mtime).strftime('%H:%M:%S')
    else:
        try:
            if ' ' in timestamp:
                timestamp = timestamp.split()[1] if ':' in timestamp.split()[1] else timestamp.split()[0]
            # 确保格式为 HH:MM:SS
            if timestamp.count(':') == 1:
                timestamp += ':00'
        except:
            timestamp = datetime.fromtimestamp(file_mtime).strftime('%H:%M:%S')

    row = {
        'File': filename,
        'Time': timestamp,
        '内核占用_KB': kernel_kb,
        '内核占用_MB': round(kernel_kb / 1024, 2),
        '驱动占用_KB': driver_kb,
        '驱动占用_MB': round(driver_kb / 1024, 2),
        'zones_KB': data['zones'],
        'wired_residual_KB': data['wired_residual'],
        'VM_KERN_总和_KB': sum(data['vm_kern'].values()),
        'VM_KERN_COUNT_BOOT_STOLEN_KB': data['vm_kern'].get('VM_KERN_COUNT_BOOT_STOLEN', 0),
        'com.apple.*_总和_KB': sum(data['drivers'].values()),
        'com.apple.filesystems.apfs_KB': data['drivers'].get('com.apple.filesystems.apfs', 0),
    }

    # 添加所有 VM_KERN_* 字段
    for key in sorted(data['vm_kern'].keys()):
        row[key] = data['vm_kern'][key]

    for key in sorted(data['drivers'].keys()):
        row[key] = data['drivers'][key]

    return row


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))

    zprint_files = []
    for fname in os.listdir(script_dir):
        if fname.startswith('zprint_') and not fname.endswith(('.py', '.csv', '.xlsx')):
            zprint_files.append(fname)

    if not zprint_files:
        print("错误: 未找到 zprint 文件")
        return

    def extract_number(filename):
        match = re.search(r'zprint_(\d+)', filename)
        return int(match.group(1)) if match else 0

    zprint_files.sort(key=extract_number)

    results = []
    for fname in zprint_files:
        filepath = os.path.join(script_dir, fname)
        print(f"正在解析: {fname}")

        file_mtime = os.path.getmtime(filepath)

        data = parse_zprint_file(filepath)
        kernel_kb, driver_kb = calculate_memory(data)
        row = format_output_row(fname, data, kernel_kb, driver_kb, file_mtime)
        results.append(row)

    # 输出到 XLSX
    if results:
        output_file = os.path.join(script_dir, 'ios_zprint_report.xlsx')

        # 合并所有文件的字段（保持固定列在前，动态列按字母排序在后）
        fixed_cols = list(results[0].keys())[:12]  # 前12个固定汇总列（增加了 Time 列）
        all_dynamic = set()
        for row in results:
            all_dynamic.update(k for k in row if k not in fixed_cols)
        fieldnames = fixed_cols + sorted(all_dynamic)

        wb = Workbook()
        ws = wb.active
        ws.title = "zprint_report"

        header_font = Font(bold=True)
        for col_idx, field in enumerate(fieldnames, 1):
            cell = ws.cell(row=1, column=col_idx, value=field)
            cell.font = header_font

        for row_idx, row_data in enumerate(results, 2):
            for col_idx, field in enumerate(fieldnames, 1):
                ws.cell(row=row_idx, column=col_idx, value=row_data.get(field, 0))

        wb.save(output_file)

        print(f"\n✓ 成功输出: {output_file}")
        print(f"✓ 共处理 {len(results)} 个文件")
        # print(f"\n公式说明:")
        # print(f"  内核占用 = (所有 VM_KERN_* + zones) - VM_KERN_COUNT_BOOT_STOLEN + WIRED_残余")
        # print(f"  驱动占用 = sum(com.apple.*) - com.apple.filesystems.apfs")
    else:
        print("错误: 没有数据可输出")


if __name__ == '__main__':
    main()
