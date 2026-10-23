#!/usr/bin/env python3
"""解析 memview 文件，计算内核占用和驱动内存，输出 Excel。"""

import glob
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    print("需要 openpyxl，正在安装...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    import openpyxl

# 需要提取的字段（kB）
COLON_FIELDS = [
    "KObject", "Pagetable", "DireAllocate", "ActvStack",
    "SUnreclaim", "VmemmapRsvTotal",
    "CmaUsed", "KDrvUsed", "UDrvUsed", "IonTotalUsed", "GpuTotalUsed",
]

PROC_FIELDS = [
    "sysmgr-main", "crypto-main", "devmgr.elf",
    "dh-actv-elastic", "hguard.elf",
]


def parse_file(filepath):
    """从单个 memview 文件中提取所需字段值（kB）。"""
    data = {}
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue

            # 带冒号的字段: "KObject:                 62208 kB"
            for key in COLON_FIELDS:
                if line.startswith(key + ":"):
                    m = re.search(r":\s+([\d]+)", line)
                    if m:
                        data[key] = int(m.group(1))
                    break

            # 进程字段: "sysmgr-main              162550 (0 in SwapPss) kB"
            for key in PROC_FIELDS:
                if line.startswith(key):
                    m = re.search(r"\s+(\d+)\s+\(", line)
                    if m:
                        data[key] = int(m.group(1))
                    break

    return data


def calc(data):
    """计算内核占用和驱动内存（kB）。"""
    kernel = (
        data.get("KObject", 0)
        + data.get("Pagetable", 0)
        + data.get("DireAllocate", 0)
        + data.get("ActvStack", 0)
        + data.get("SUnreclaim", 0)
        + data.get("VmemmapRsvTotal", 0)
        + data.get("sysmgr-main", 0)
        + data.get("dh-actv-elastic", 0)
        + data.get("crypto-main", 0)
        + data.get("devmgr.elf", 0)
        + data.get("hguard.elf", 0)
    )
    driver = (
        data.get("CmaUsed", 0)
        + data.get("KDrvUsed", 0)
        + data.get("UDrvUsed", 0)
        - data.get("IonTotalUsed", 0)
        - data.get("GpuTotalUsed", 0)
    )
    return kernel, driver


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
    pattern = os.path.join(folder, "memview_*_PerformanceDynamic_*")
    files = sorted(glob.glob(pattern))

    if not files:
        print(f"在 {folder} 中未找到 memview 文件")
        return

    # 所有列名
    all_fields = COLON_FIELDS + PROC_FIELDS
    col_headers = [f + " (kB)" for f in all_fields] + ["内核占用 (MB)", "驱动内存 (MB)"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "内存分析"

    # 写表头
    ws.cell(row=1, column=1, value="文件名")
    for ci, h in enumerate(col_headers, start=2):
        ws.cell(row=1, column=ci, value=h)

    # 逐文件处理
    for ri, fpath in enumerate(files, start=2):
        fname = os.path.basename(fpath)
        data = parse_file(fpath)
        kernel_kb, driver_kb = calc(data)

        ws.cell(row=ri, column=1, value=fname)
        for ci, key in enumerate(all_fields, start=2):
            ws.cell(row=ri, column=ci, value=data.get(key, 0))
        ws.cell(row=ri, column=len(all_fields) + 2, value=round(kernel_kb / 1024, 2))
        ws.cell(row=ri, column=len(all_fields) + 3, value=round(driver_kb / 1024, 2))

    out_path = os.path.join(folder, "mem_analysis.xlsx")
    wb.save(out_path)
    print(f"已生成: {out_path}  ({len(files)} 个文件)")


if __name__ == "__main__":
    main()
