#!/usr/bin/env python3
"""将 all_log.txt 中的步骤标记按时间对齐到 ios_zprint_report.xlsx"""
import re, os, sys
import openpyxl
from collections import OrderedDict

LOG_FILE  = 'all_log.txt'
XLSX_FILE = 'ios_zprint_report.xlsx'
OUTPUT    = 'ios_zprint_report_with_steps.xlsx'

def hms_to_sec(t):
    """HH:MM:SS -> 秒数"""
    h, m, s = t.split(':')
    return int(h)*3600 + int(m)*60 + int(s)

def parse_log(path):
    """解析日志，提取 (秒数, HH:MM:SS, APP名, 步骤名)"""
    results = []
    # iTrace: trace 同步日志内容 [APP名][步骤名]
    p1 = re.compile(
        r'(\d{2}:\d{2}:\d{2}),\d+ - iTrace\.py.*?INFO: trace 同步日志内容'
        r" \[([^\]]+)\]\[([^\]]+)\]"
    )
    # PerformanceDynamic_XXX_00XX.py ... INFO: 步骤名
    p2 = re.compile(
        r'(\d{2}:\d{2}:\d{2}),\d+ - PerformanceDynamic_(\w+?)_\d+\.py'
        r'.*?INFO:\s*(.+)'
    )
    skip = {'测试环境开始准备','用例开始执行','用例执行结束'}

    with open(path, 'r', encoding='utf-8') as f:
        for line in f:
            m = p1.search(line)
            if m:
                t, app, step = m.group(1), m.group(2), m.group(3)
                step = step.strip("'\" ")
                results.append((hms_to_sec(t), t, app, step))
                continue
            m = p2.search(line)
            if m:
                t, app, step = m.group(1), m.group(2), m.group(3).strip()
                if step in skip or step.startswith('等待'):
                    continue
                results.append((hms_to_sec(t), t, app, step))
    return results


def main():
    steps = parse_log(LOG_FILE)
    # 每个 (app, step) 只保留最早出现的时间
    seen = OrderedDict()
    for secs, t, app, step in steps:
        key = f"{app}-{step}"
        if key not in seen or secs < seen[key][0]:
            seen[key] = (secs, t)

    # 读取 xlsx
    wb = openpyxl.load_workbook(XLSX_FILE)
    ws = wb.active
    # 收集每行的秒数
    row_secs = []
    for r in range(2, ws.max_row + 1):
        tv = ws.cell(r, 2).value
        row_secs.append((r, hms_to_sec(str(tv)) if tv else None))

    # 在第3列插入新列
    ws.insert_cols(3)
    ws.cell(1, 3, '步骤标记')

    # 对每个步骤，找 <= 步骤时间 且最接近的行
    for label, (s, t) in seen.items():
        best_row, best_diff = None, float('inf')
        for r, rs in row_secs:
            if rs is None:
                continue
            d = s - rs
            if 0 <= d < best_diff:
                best_diff = d
                best_row = r
        if best_row:
            old = ws.cell(best_row, 3).value
            ws.cell(best_row, 3, f"{old}\n{label}" if old else label)

    wb.save(OUTPUT)
    print(f"完成! 输出: {OUTPUT}")
    print(f"共提取 {len(seen)} 个步骤:")
    for label, (s, t) in seen.items():
        print(f"  {t}  {label}")

if __name__ == '__main__':
    main()
