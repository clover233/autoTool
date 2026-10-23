import os, re
from openpyxl import Workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

# 指定compressor文件路径
COMPRESSOR_FILE = "/Users/clover/Result0420/compressor/compressor5"
OUTPUT_DIR = "/Users/clover/Result0420/compressor"
output_file = os.path.join(OUTPUT_DIR, 'compressor5.xlsx')

print(f"Reading {COMPRESSOR_FILE}...")

# 解析数据
data_rows = []

with open(COMPRESSOR_FILE, 'r', encoding='utf-8', errors='replace') as f:
    for line in f:
        line = line.strip()

        # 跳过命令行和空行
        if not line or line.startswith('iPhone14:'):
            continue

        # 解析格式: ts: 2026-04-20 09:24:11, compact awakens: 82, compress size: 228MB, ...
        if line.startswith('ts:'):
            # 提取时间戳
            ts_match = re.search(r'ts:\s*([^,]+)', line)
            timestamp = ts_match.group(1).strip() if ts_match else ''

            # 提取 compact awakens
            awakens_match = re.search(r'compact awakens:\s*(\d+)', line)
            compact_awakens = awakens_match.group(1) if awakens_match else ''

            # 提取 compress size
            compress_match = re.search(r'compress size:\s*([^,]+)', line)
            compress_size = compress_match.group(1).strip() if compress_match else ''

            # 提取 pool size
            pool_match = re.search(r'pool size:\s*([^,]+)', line)
            pool_size = pool_match.group(1).strip() if pool_match else ''

            # 提取 pool actual size
            actual_match = re.search(r'pool actual size:\s*([^,]+)', line)
            pool_actual_size = actual_match.group(1).strip() if actual_match else ''

            # 提取 minor frag
            minor_match = re.search(r'minor frag:\s*([\d.]+)', line)
            minor_frag = minor_match.group(1) if minor_match else ''

            # 提取 major frag
            major_match = re.search(r'major frag:\s*([\d.]+)', line)
            major_frag = major_match.group(1) if major_match else ''

            data_rows.append({
                'timestamp': timestamp,
                'compact_awakens': compact_awakens,
                'compress_size': compress_size,
                'pool_size': pool_size,
                'pool_actual_size': pool_actual_size,
                'minor_frag': minor_frag,
                'major_frag': major_frag
            })

print(f"Found {len(data_rows)} data rows")

# 创建 Excel 工作簿
wb = Workbook()
ws = wb.active
ws.title = 'compressor_data'

# 构建表头
headers = ['timestamp', 'compact_awakens', 'compress_size', 'pool_size',
           'pool_actual_size', 'minor_frag', 'major_frag']
ws.append(headers)

# 写入数据
for row_data in data_rows:
    row = [
        row_data['timestamp'],
        row_data['compact_awakens'],
        row_data['compress_size'],
        row_data['pool_size'],
        row_data['pool_actual_size'],
        row_data['minor_frag'],
        row_data['major_frag']
    ]
    ws.append(row)

# 添加网格线
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.border = thin_border

# 自动调整列宽
for col_idx in range(1, ws.max_column + 1):
    max_length = 0
    column_letter = get_column_letter(col_idx)

    for row_idx in range(1, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        try:
            if cell.value:
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length
        except:
            pass

    adjusted_width = min(max_length + 2, 50)
    ws.column_dimensions[column_letter].width = adjusted_width

# 保存文件
wb.save(output_file)
print(f"Done! {len(data_rows)} rows -> {output_file}")
