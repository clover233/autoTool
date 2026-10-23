#!/usr/bin/env python3
"""
Parse footprint JSON -> xlsx (single sheet with all processes).
Each process block has 8 rows (title + header + 6 metrics) + 1 blank separator.
"""

import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def fmt_bytes(n):
    if n == 0:
        return "0 B"
    abs_n = abs(n)
    if abs_n >= 1024 * 1024:
        val = n / (1024 * 1024)
        s = f"{val:.1f}".rstrip('0').rstrip('.')
        return f"{s} MB"
    elif abs_n >= 1024:
        val = n / 1024
        s = f"{val:.1f}".rstrip('0').rstrip('.')
        return f"{s} KB"
    else:
        return f"{n} B"


BLUE_FILL   = PatternFill("solid", fgColor="4472C4")
GRAY_FILL   = PatternFill("solid", fgColor="D9E1F2")
ALT_FILL    = PatternFill("solid", fgColor="EEF2FF")
TITLE_FILL  = PatternFill("solid", fgColor="FFC000")

HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT  = Font(bold=True, size=11)
LABEL_FONT  = Font(bold=True, size=10)

thin = Side(style="thin")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")
LEFT   = Alignment(horizontal="left",   vertical="center")

METRICS = [
    ("Dirty",       "dirty"),
    ("(Swapped)",   "swapped"),
    ("Clean",       "clean"),
    ("Reclaimable", "reclaimable"),
    ("(Wired)",     "wired"),
    ("Regions",     "regions"),
]


def main():
    input_path  = sys.argv[1] if len(sys.argv) > 1 else "footprint_35.json"
    output_path = sys.argv[2] if len(sys.argv) > 2 else "footprint_35.xlsx"

    print(f"Reading {input_path} ...")
    with open(input_path, encoding="utf-8") as f:
        data = json.load(f)

    processes = data.get("processes", [])
    print(f"Found {len(processes)} processes")

    wb = Workbook()
    ws = wb.active
    ws.title = "All Processes"

    current_row = 1
    max_cols = 1

    for proc_idx, proc in enumerate(processes):
        name      = proc.get("name", "unknown")
        pid       = proc.get("pid", 0)
        page_size = proc.get("page size", 16384)
        footprint = proc.get("footprint", 0)
        categories = proc.get("categories", {})

        # skip if has errors or no footprint
        if "errors" in proc or footprint == 0:
            continue

        # sort by dirty desc
        sorted_cats = sorted(categories.items(),
                             key=lambda x: x[1].get("dirty", 0), reverse=True)
        cat_names = [c[0] for c in sorted_cats]
        n_cats = len(cat_names)
        max_cols = max(max_cols, 1 + n_cats)

        # ── Row: title ────────────────────────────────────────────────────────
        fp_mb = footprint / (1024 * 1024)
        title = f"{name} [{pid}]    Footprint: {fp_mb:.0f} MB  ({page_size} bytes per page)"

        if n_cats > 0:
            ws.merge_cells(start_row=current_row, start_column=1,
                           end_row=current_row,   end_column=1 + n_cats)
        c = ws.cell(row=current_row, column=1, value=title)
        c.font = TITLE_FONT
        c.fill = TITLE_FILL
        c.alignment = LEFT
        ws.row_dimensions[current_row].height = 20
        current_row += 1

        # ── Row: category header ──────────────────────────────────────────────
        ws.cell(row=current_row, column=1, value="").border = BORDER

        for col, cat in enumerate(cat_names, start=2):
            c = ws.cell(row=current_row, column=col, value=cat)
            c.font = HEADER_FONT
            c.fill = BLUE_FILL
            c.alignment = CENTER
            c.border = BORDER
        ws.row_dimensions[current_row].height = 30
        current_row += 1

        # ── Rows: metrics ─────────────────────────────────────────────────────
        for row_off, (label, key) in enumerate(METRICS):
            fill = ALT_FILL if row_off % 2 == 0 else None

            c = ws.cell(row=current_row, column=1, value=label)
            c.font = LABEL_FONT
            c.fill = GRAY_FILL
            c.alignment = LEFT
            c.border = BORDER

            for col, (cat, metrics) in enumerate(sorted_cats, start=2):
                val = metrics.get(key, 0)
                display = str(val) if key == "regions" else fmt_bytes(val)
                c = ws.cell(row=current_row, column=col, value=display)
                c.alignment = RIGHT
                c.border = BORDER
                if fill:
                    c.fill = fill

            current_row += 1

        # ── blank separator ───────────────────────────────────────────────────
        current_row += 1

        if (proc_idx + 1) % 50 == 0:
            print(f"  {proc_idx + 1}/{len(processes)} processes done ...")

    # ── column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 14
    for col in range(2, max_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    ws.freeze_panes = "B1"

    wb.save(output_path)
    print(f"Saved -> {output_path}")


if __name__ == "__main__":
    main()
