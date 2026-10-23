#!/usr/bin/env python3
"""解析 vmstats.log，每条采样写入 Excel 的一行。

依赖：python3 -m pip install openpyxl
用法：python3 parse_vmstats.py [日志路径] [-o 输出.xlsx]
不传参数时读取 DEFAULT_INPUT，默认输出到日志同目录。
"""

import argparse
import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path

DEFAULT_INPUT = Path('/Users/clover/Desktop/SeaOfStars7i/Result/result1/vmstats.log')
LINE_PATTERN = re.compile(r'^\[(\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})\]\s*(\{.*\})$')


def flatten_fields(value, prefix='', result=None):
    """将嵌套 JSON 展开为带完整路径的字段，保留原始数值。"""
    if result is None:
        result = {}
    if isinstance(value, dict):
        for key, child in value.items():
            flatten_fields(child, f'{prefix}.{key}' if prefix else key, result)
    else:
        if prefix in result:
            raise ValueError(f'重复字段路径：{prefix}')
        result[prefix] = json.dumps(value, ensure_ascii=False) if isinstance(value, list) else value
    return result


def parse_log(input_path):
    """按日志顺序保留记录，汇总所有字段；错误包含源文件行号。"""
    fields = {}
    records = []
    with Path(input_path).open(encoding='utf-8-sig') as source:
        for line_number, line in enumerate(source, start=1):
            if not line.strip():
                continue
            try:
                match = LINE_PATTERN.fullmatch(line.strip())
                if not match:
                    raise ValueError('应为 [yyyy-MM-dd HH:mm:ss] JSON')
                timestamp = datetime.strptime(match[1], '%Y-%m-%d %H:%M:%S')
                values = flatten_fields(json.loads(match[2]))
                values.pop('sessionId', None)
                if 'timestamp' in values:
                    raise ValueError('JSON 字段 timestamp 与时间列重名')
                fields.update(dict.fromkeys(values))
                records.append((timestamp, values))
            except (ValueError, TypeError) as exc:
                raise ValueError(f'第 {line_number} 行解析失败：{exc}') from exc
    if not records:
        raise ValueError('日志中没有有效记录')
    if len(records) > 1_048_574 or len(fields) + 1 > 16_384:
        raise ValueError('数据超过 Excel 单表容量')
    return ['timestamp', *fields], records


def convert(input_path, output_path=None):
    """导出 xlsx；时间为 Excel 日期，缺失字段留空，不合并重复时间。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    input_path = Path(input_path).expanduser().resolve()
    output_path = Path(output_path).expanduser().resolve() if output_path else input_path.with_suffix('.xlsx')
    if input_path == output_path:
        raise ValueError('输出不能覆盖源日志')
    if output_path.suffix.lower() != '.xlsx':
        raise ValueError('输出文件必须使用 .xlsx 扩展名')
    headers, records = parse_log(input_path)
    # 按父级分组，保证后续采样新增的字段也落在对应合并表头下。
    groups = {}
    for key in headers[1:]:
        parts = key.split('.')
        if parts[0] == 'value':
            parts = parts[1:]
        group = '.'.join(parts[:-1])
        groups.setdefault(group, []).append(key)
    headers = ['timestamp', *(key for keys in groups.values() for key in keys)]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'vmstats'
    sheet.append([None] * len(headers))
    sheet.append([key.rsplit('.', 1)[-1] for key in headers])
    column = 2
    for group, keys in groups.items():
        if group:
            sheet.cell(1, column, group)
            if len(keys) > 1:
                sheet.merge_cells(start_row=1, start_column=column,
                                  end_row=1, end_column=column + len(keys) - 1)
            sheet.cell(1, column).alignment = Alignment(horizontal='center')
        column += len(keys)
    for timestamp, values in records:
        sheet.append([timestamp, *(values.get(key) for key in headers[1:])])
    for row in sheet.iter_rows():
        for cell in row:
            # JSON 字符串始终作为文本，不能被 Excel 解释成公式。
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for row in sheet.iter_rows(min_row=3):
        row[0].number_format = 'yyyy-mm-dd hh:mm:ss'
    # openpyxl 不执行 Excel 的 AutoFit，按实际显示内容估算列宽。
    # 中文等宽字符计为两个字符；合并分组标题不参与单列宽度计算。
    for index in range(1, len(headers) + 1):
        width = 0
        for cells in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
            value = cells[0].value
            text = value.strftime('%Y-%m-%d %H:%M:%S') if isinstance(value, datetime) else str(value) if value is not None else ''
            width = max(width, sum(2 if unicodedata.east_asian_width(char) in ('W', 'F') else 1 for char in text))
        sheet.column_dimensions[get_column_letter(index)].width = min(255, width + 3)
    sheet.freeze_panes = 'B3'
    sheet.auto_filter.ref = f'A2:{get_column_letter(len(headers))}{sheet.max_row}'
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path, len(records), len(headers)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('input', nargs='?', type=Path, default=DEFAULT_INPUT, help='日志路径')
    parser.add_argument('-o', '--output', type=Path, help='输出路径，默认日志同目录、同名 .xlsx')
    args = parser.parse_args()
    try:
        output, rows, columns = convert(args.input, args.output)
    except ModuleNotFoundError as exc:
        if exc.name != 'openpyxl':
            raise
        parser.exit(1, '缺少依赖，请运行：python3 -m pip install openpyxl\n')
    except (OSError, ValueError) as exc:
        parser.exit(1, f'转换失败：{exc}\n')
    print(f'已导出 {rows} 条记录、{columns} 列：{output}')


if __name__ == '__main__':
    main()
