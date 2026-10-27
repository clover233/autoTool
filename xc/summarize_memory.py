#!/usr/bin/env python3
"""按录制时间排序，提取每个 trace XML 的第 2 个采样点并生成内存汇总。

安装依赖：python3 -m pip install openpyxl
运行：python3 summarize_memory.py --result-dir /你的/Result/result1
仅校验：python3 summarize_memory.py --result-dir /你的/Result/result1 --check-only

读取已由 export_activity_monitor.py 导出的 XML，不直接读取 .trace 包。
采样点按 start 数值排序；第 2 点是第 2 个不同时间戳，不是第 2 行进程数据。
内存只取 memory-physical-footprint；MiB = bytes / 1024 / 1024。
"""

import argparse
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET

from xctrace_tools.config import (ACTIVITY_SCHEMA, APPS, LOAD_APPS, CATEGORIES,
                                 DEFAULT_RESULT_DIR, app_for_process, classify_process)
from xctrace_tools.files import find_activity_xml as find_xml_files
from xctrace_tools.logs import step_at_sample as read_step_log
from xctrace_tools.xmlio import TraceTable, process_name


def read_second_sample(path):
    table = TraceTable(path, ACTIVITY_SCHEMA)
    table.require('start', 'process', 'memory-physical-footprint')
    rows = list(table.rows())
    timestamps = sorted({int(table.resolve(row['start']).text) for row in rows})
    if len(timestamps) < 2:
        raise ValueError(f'只有 {len(timestamps)} 个采样点，无法读取第 2 个')
    if timestamps[0] != 0:
        raise ValueError(f'第一个采样点不是 0 秒（{timestamps[0]} ns），请检查 trace')
    selected = timestamps[1]
    records = []
    for row in rows:
        if int(table.resolve(row['start']).text) != selected:
            continue
        process = process_name(table.resolve(row['process']))
        if not process:
            raise ValueError('进程名为空')
        footprint = int(table.resolve(row['memory-physical-footprint']).text)
        if footprint < 0:
            raise ValueError(f'进程 {process} 的 footprint 为负数')
        pid = int(table.resolve(row['pid']).text) if 'pid' in row else None
        records.append({'process': process, 'pid': pid, 'bytes': footprint})
    return selected, len(timestamps), records


def trace_start_time(path):
    # ...-step_4-200022-20260929_200026(4s).xml：前一个时间是录制开始时间。
    match = re.search(r'-step_\d+-(\d{6})-(\d{8})_(\d{6})', path.stem)
    if match is None:
        raise ValueError('文件名缺少录制时间，无法按时间排序')
    start = datetime.strptime(match[2] + match[1], '%Y%m%d%H%M%S')
    end = datetime.strptime(match[2] + match[3], '%Y%m%d%H%M%S')
    if start > end:  # 跨午夜时文件名中的日期对应结束时间。
        start -= timedelta(days=1)
    return start


def collect(result_dir):
    result_dir = Path(result_dir).expanduser()
    errors = []
    dated = []
    for path in find_xml_files(result_dir):
        try:
            dated.append((trace_start_time(path), path))
        except (ValueError, OSError) as exc:
            errors.append((str(path.relative_to(result_dir)), str(exc)))
    dated.sort(key=lambda pair: (pair[0], str(pair[1])))
    traces = []
    for start, path in dated:
        try:
            sample_ns, sample_count, records = read_second_sample(path)
            foreground, action, note = read_step_log(path, sample_ns)
            configured = foreground in set(APPS.values())
            if not foreground:
                status = '前台未知'
            elif foreground == '负载':
                status = '负载阶段'
            elif not configured:
                status = '前台未配置'
            else:
                status = '已匹配'
            totals = Counter({cat: 0 for cat in CATEGORIES})
            counts = Counter()
            for record in records:
                record['category'] = classify_process(record['process'], foreground)
                totals[record['category']] += record['bytes']
                counts[record['category']] += 1
            if status == '已匹配' and counts['ForegroundAPP'] == 0:
                status = '未采集到前台进程'
            assert sum(totals.values()) == sum(record['bytes'] for record in records)
            traces.append({'file': str(path.relative_to(result_dir)), 'start': start,
                           'sample_ns': sample_ns, 'sample_count': sample_count,
                           'step': int(re.search(r'-step_(\d+)-', path.stem)[1]),
                           'app': foreground, 'action': action, 'note': note,
                           'status': status, 'records': records, 'totals': totals})
        except (ValueError, OSError, ET.ParseError, TypeError, KeyError, IndexError) as exc:
            errors.append((str(path.relative_to(result_dir)), str(exc)))
    return traces, errors


def export_excel(traces, errors, output):
    output = Path(output).expanduser()
    if output.suffix.lower() != '.xlsx':
        raise ValueError('输出文件扩展名必须为 .xlsx')
    if any(not trace['records'] for trace in traces):
        raise ValueError('无法汇总空采样点')
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise SystemExit('缺少依赖，请运行：python3 -m pip install openpyxl')

    book = Workbook()
    summary = book.active
    summary.title = 'trace内存汇总'
    detail = book.create_sheet('第2采样点进程明细')
    rules = book.create_sheet('分类规则')
    summary.append(['序号', 'trace开始时间', '步骤id', '前台应用', '步骤',
                    '第2采样点(秒)', *[f'{cat} (MiB)' for cat in CATEGORIES],
                    '总计 (MiB)', '进程数', '前台匹配状态'])
    detail.append(['序号', 'trace开始时间', '步骤id', '前台应用', '步骤',
                   '第2采样点(秒)', 'proc_name', '应用名', 'proc_type', 'pid',
                   'physical memory footprint (bytes)', 'physical memory footprint (MiB)'])
    for trace_id, trace in enumerate(traces, 1):
        start_detail = detail.max_row + 1
        for record in trace['records']:
            detail.append([trace_id, trace['start'], trace['step'],
                           trace['app'], trace['action'], trace['sample_ns'] / 1e9,
                           record['process'], app_for_process(record['process']),
                           record['category'], record['pid'], record['bytes'],
                           f'=K{detail.max_row + 1}/1048576'])
        end_detail = detail.max_row
        prefix = [trace_id, trace['start'], trace['step'], trace['app'],
                  trace['action'], trace['sample_ns'] / 1e9]
        category_formulas = [
            f'=SUMIFS(\'第2采样点进程明细\'!$L${start_detail}:$L${end_detail},'
            f'\'第2采样点进程明细\'!$I${start_detail}:$I${end_detail},"{cat}")'
            for cat in CATEGORIES
        ]
        if trace['status'] in ('前台未知', '前台未配置'):
            category_formulas[0] = 'n.a.'
        # 总计直接汇总原始记录，避免 n.a. 或任何类别遗漏导致总计偏小。
        total = f'=SUM(\'第2采样点进程明细\'!L{start_detail}:L{end_detail})'
        summary.append(prefix + category_formulas + [total, len(trace['records']),
                                                     trace['status']])

    rules.append(['类型', '进程名', '中文应用名'])
    for name, mapping in [('ForegroundAPP / BackgroudAPP', APPS), ('BackgroudLoad', LOAD_APPS)]:
        for key, value in mapping.items():
            rules.append([name, key, value])
    rules.append(['SystemService', '其余进程', '按清单分类，未配置应用也可能归入系统服务'])
    rules.append(['内存单位', 'MiB', 'memory-physical-footprint bytes / 1048576'])
    rules.append(['取样规则', '第 2 个不同的 start 时间戳', '首点必须是 0；没有第 2 点则记录失败'])
    rules.append(['前台判断', '配套同名 .log', '日志前采样按该 trace 首条步骤标注'])
    rules.append(['普通应用', '应用名与前台相同：ForegroundAPP', '其余普通应用：BackgroudAPP'])
    rules.append(['负载优先', 'LOAD_APPS 始终为 BackgroudLoad', '进程名和中文应用名精确匹配，区分大小写'])
    rules.append(['未配置前台', 'ForegroundAPP 显示 n.a.', '请在 xctrace_tools/config.py 的 APPS 字典补充相应进程'])
    if errors:
        failure = book.create_sheet('解析失败')
        failure.append(['来源XML', '失败原因'])
        for error in errors:
            failure.append(error)

    widths = {
        summary.title: [8, 23, 10, 20, 55, 18, 22, 22, 22, 22, 20, 12, 25],
        detail.title: [8, 23, 10, 20, 55, 18, 36, 22, 23, 12, 34, 34],
        rules.title: [34, 45, 70],
        '解析失败': [80, 70],
    }
    for sheet in book:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'G2' if sheet in (summary, detail) else 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 36
        for column, width in enumerate(widths[sheet.title], 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        for cell in sheet[1]:
            cell.font = Font(name='Arial', bold=True, color='1F3A53', size=10)
            cell.fill = PatternFill('solid', fgColor='DDEBF7')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for cells in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[cells[0].row].height = 32 if sheet in (summary, detail) else 24
            for cell in cells:
                cell.font = Font(name='Arial', size=10)
                cell.alignment = Alignment(vertical='center', wrap_text=cell.column == (5 if sheet in (summary, detail) else 3))
                if isinstance(cell.value, datetime):
                    cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            if sheet in (summary, detail):
                sheet.cell(cells[0].row, 6).number_format = '0.000000000'
        if sheet == summary:
            for cells in sheet.iter_rows(min_row=2, min_col=7, max_col=11):
                for cell in cells:
                    cell.number_format = '#,##0.00'
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 13).value != '已匹配':
                    sheet.cell(row, 13).fill = PatternFill('solid', fgColor='FFF0D6')
        elif sheet == detail:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, 12).number_format = '#,##0.00'
    for sheet in book:
        text_columns = ({4, 5, 13} if sheet == summary else {4, 5, 7, 8, 9} if sheet == detail
                        else set(range(1, sheet.max_column + 1)))
        for cells in sheet:
            for cell in cells:
                if cell.column in text_columns and isinstance(cell.value, str):
                    cell.data_type = 's'
    # Excel/WPS 打开时重算公式；脚本校验不依赖未生成的公式缓存。
    from openpyxl.workbook.properties import CalcProperties
    book.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--result-dir', '-r', type=Path, default=Path(DEFAULT_RESULT_DIR))
    parser.add_argument('--output', '-o', type=Path, help='默认输出到 result 目录的 physical_memory_summary.xlsx；同名结果会替换')
    parser.add_argument('--check-only', action='store_true', help='只解析和校验，不生成 Excel')
    args = parser.parse_args(argv)
    result_dir = args.result_dir.expanduser().resolve()
    if not result_dir.is_dir():
        parser.error(f'结果目录不存在: {result_dir}')
    traces, errors = collect(result_dir)
    if not traces and not errors:
        parser.error('没有找到 PerformanceDynamic_*.xml；请先导出 trace XML')
    record_count = sum(len(trace['records']) for trace in traces)
    print(f'成功 trace: {len(traces)}；失败: {len(errors)}；第2采样点进程记录: {record_count}')
    for file, error in errors:
        print(f'失败: {file}: {error}', file=sys.stderr)
    unmapped = sorted({trace['app'] for trace in traces if trace['status'] == '前台未配置'})
    if unmapped:
        print('未配置前台应用（可在 xctrace_tools/config.py 的 APPS 字典补充）：' + '、'.join(unmapped))
    if traces:
        print(f'时间范围: {traces[0]["start"]} 至 {traces[-1]["start"]}')
        print('第2采样点示例（秒）: ' + ', '.join(f'{trace["sample_ns"] / 1e9:.9f}' for trace in traces[:3]))
    if not args.check_only:
        output = (args.output or result_dir / 'physical_memory_summary.xlsx').expanduser().resolve()
        if output.suffix.lower() != '.xlsx':
            parser.error('输出文件扩展名必须为 .xlsx')
        try:
            export_excel(traces, errors, output)
        except (OSError, ValueError) as exc:
            parser.exit(1, f'导出失败：{exc}\n')
        print(f'已输出: {output}')
    return 1 if errors else 0


if __name__ == '__main__':
    sys.exit(main())
